from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import HTMLResponse
from sqlalchemy.orm import Session
from .. import models
from ..database import get_db

router = APIRouter()

@router.get("/{run_id}/report", response_class=HTMLResponse)
def get_execution_report(run_id: int, db: Session = Depends(get_db)):
    run = db.query(models.ExecutionRun).filter(models.ExecutionRun.id == run_id).first()
    if not run:
        raise HTTPException(status_code=404, detail="Run not found")
        
    script = db.query(models.TestScript).filter(models.TestScript.id == run.script_id).first()
    
    # Calculate duration
    duration = "Running"
    if run.completed_at and run.started_at:
        duration = f"{(run.completed_at - run.started_at).total_seconds():.1f}s"
        
    html = f"""
    <!DOCTYPE html>
    <html lang="en">
    <head>
        <meta charset="UTF-8">
        <meta name="viewport" content="width=device-width, initial-scale=1.0">
        <title>Execution Report - Run {run.id}</title>
        <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap" rel="stylesheet">
        <style>
            :root {{
                --bg-main: #0f172a;
                --bg-card: #1e293b;
                --text-main: #f8fafc;
                --text-muted: #94a3b8;
                --border-color: #334155;
                --accent-primary: #f97316;
                --success: #22c55e;
                --error: #ef4444;
            }}
            body {{
                font-family: 'Inter', sans-serif;
                background-color: var(--bg-main);
                color: var(--text-main);
                margin: 0;
                padding: 2rem;
                line-height: 1.5;
            }}
            .container {{
                max-width: 1000px;
                margin: 0 auto;
            }}
            .header-card {{
                background: var(--bg-card);
                border: 1px solid var(--border-color);
                border-radius: 12px;
                padding: 2rem;
                margin-bottom: 2rem;
                box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.1);
            }}
            .title-area {{
                display: flex;
                justify-content: space-between;
                align-items: flex-start;
                border-bottom: 1px solid var(--border-color);
                padding-bottom: 1.5rem;
                margin-bottom: 1.5rem;
            }}
            h1 {{
                margin: 0;
                font-size: 1.8rem;
                font-weight: 600;
            }}
            .meta-grid {{
                display: grid;
                grid-template-columns: repeat(auto-fit, minmax(200px, 1fr));
                gap: 1.5rem;
            }}
            .meta-item label {{
                display: block;
                color: var(--text-muted);
                font-size: 0.85rem;
                text-transform: uppercase;
                letter-spacing: 0.05em;
                margin-bottom: 0.5rem;
            }}
            .meta-item .value {{
                font-size: 1.1rem;
                font-weight: 500;
            }}
            .badge {{
                display: inline-flex;
                align-items: center;
                padding: 0.25rem 0.75rem;
                border-radius: 9999px;
                font-size: 0.875rem;
                font-weight: 600;
            }}
            .badge.passed {{ background-color: rgba(34, 197, 94, 0.2); color: #4ade80; border: 1px solid rgba(34, 197, 94, 0.3); }}
            .badge.failed {{ background-color: rgba(239, 68, 68, 0.2); color: #f87171; border: 1px solid rgba(239, 68, 68, 0.3); }}
            .badge.running {{ background-color: rgba(234, 179, 8, 0.2); color: #fde047; border: 1px solid rgba(234, 179, 8, 0.3); }}
            
            .assets-card {{
                background: var(--bg-card);
                border: 1px solid var(--border-color);
                border-radius: 12px;
                padding: 1.5rem;
                margin-bottom: 2rem;
            }}
            .assets-card h3 {{ margin-top: 0; margin-bottom: 1rem; font-size: 1.2rem; }}
            .btn {{
                display: inline-block;
                background-color: var(--accent-primary);
                color: white;
                padding: 0.5rem 1rem;
                text-decoration: none;
                border-radius: 6px;
                font-weight: 500;
                font-size: 0.9rem;
                transition: background-color 0.2s;
            }}
            .btn:hover {{ background-color: #ea580c; }}
            video {{
                width: 100%;
                border-radius: 8px;
                margin-top: 1rem;
                border: 1px solid var(--border-color);
            }}
            
            .timeline {{
                display: flex;
                flex-direction: column;
                gap: 1rem;
            }}
            .step-card {{
                background: var(--bg-card);
                border: 1px solid var(--border-color);
                border-radius: 8px;
                padding: 1.25rem;
                display: flex;
                gap: 1.5rem;
            }}
            .step-number {{
                background: #334155;
                color: white;
                width: 32px;
                height: 32px;
                border-radius: 50%;
                display: flex;
                align-items: center;
                justify-content: center;
                font-weight: bold;
                flex-shrink: 0;
            }}
            .step-content {{ flex: 1; }}
            .step-header {{
                display: flex;
                justify-content: space-between;
                margin-bottom: 0.5rem;
            }}
            .step-action {{
                font-weight: 600;
                font-size: 1.05rem;
            }}
            .screenshot-thumb {{
                max-width: 300px;
                border-radius: 6px;
                border: 1px solid var(--border-color);
                margin-top: 1rem;
                cursor: pointer;
                transition: transform 0.2s;
            }}
            .screenshot-thumb:hover {{
                transform: scale(1.02);
            }}
            .error-box {{
                background: rgba(239, 68, 68, 0.1);
                border: 1px solid rgba(239, 68, 68, 0.3);
                border-radius: 6px;
                padding: 1rem;
                margin-top: 1rem;
                color: #fca5a5;
                font-family: monospace;
                white-space: pre-wrap;
            }}
        </style>
    </head>
    <body>
        <div class="container">
            <div class="header-card">
                <div class="title-area">
                    <div>
                        <h1>Test Execution Report</h1>
                        <p style="color: var(--text-muted); margin: 0.5rem 0 0 0;">{script.name} (Run #{run.id})</p>
                    </div>
                    <span class="badge {run.status.lower() if run.status else ''}">{run.status}</span>
                </div>
                
                <div class="meta-grid">
                    <div class="meta-item">
                        <label>Environment</label>
                        <div class="value">{run.environment.name if run.environment else 'Unknown'}</div>
                    </div>
                    <div class="meta-item">
                        <label>Data Profile</label>
                        <div class="value">{run.data_profile.name if run.data_profile else 'None'}</div>
                    </div>
                    <div class="meta-item">
                        <label>Started At</label>
                        <div class="value">{run.started_at.strftime("%Y-%m-%d %H:%M:%S") if run.started_at else '-'}</div>
                    </div>
                    <div class="meta-item">
                        <label>Duration</label>
                        <div class="value">{duration}</div>
                    </div>
                </div>
            </div>
    """

    has_video = run.video_path is not None
    has_trace = run.trace_path is not None
    
    if has_video or has_trace:
        html += f"""
            <div class="assets-card">
                <h3>Execution Assets</h3>
                <div style="display: flex; gap: 1rem; flex-wrap: wrap;">
        """
        if has_trace:
            trace_url = "/" + run.trace_path.replace("\\", "/") if not run.trace_path.startswith("/") else run.trace_path
            if "data/" in trace_url: trace_url = "/data/" + trace_url.split("data/")[-1]
            html += f'<a href="{trace_url}" class="btn" download>↓ Download Playwright Trace</a>'
            
        if has_video:
            video_url = "/" + run.video_path.replace("\\", "/") if not run.video_path.startswith("/") else run.video_path
            if "data/" in video_url: video_url = "/data/" + video_url.split("data/")[-1]
            html += f"""
                <div style="width: 100%; margin-top: 1rem;">
                    <label style="color: var(--text-muted); font-size: 0.85rem; text-transform: uppercase;">Session Recording</label>
                    <video controls src="{video_url}"></video>
                </div>
            """
            
        html += """
                </div>
            </div>
        """

    html += """
            <h3 style="margin-top: 2rem; margin-bottom: 1.5rem;">Step Details</h3>
            <div class="timeline">
    """
    
    for step_res in sorted(run.step_results, key=lambda s: s.step.step_number if s.step else 0):
        step_number = step_res.step.step_number if step_res.step else "?"
        action = step_res.step.action if step_res.step else "Raw Execution"
        target = ""
        desc = ""
        if step_res.step:
            target = step_res.step.target_name or step_res.step.locator_value or ""
            desc = step_res.step.business_description or ""
        status_class = step_res.status.lower() if step_res.status else ""
        
        html += f"""
                <div class="step-card">
                    <div class="step-number">{step_number}</div>
                    <div class="step-content">
                        <div class="step-header">
                            <div class="step-action">{action} <span style="font-weight: 400; color: var(--text-muted);">{target}</span></div>
                            <span class="badge {status_class}" style="font-size: 0.75rem;">{step_res.status}</span>
                        </div>
                        {f'<div style="color: var(--text-muted); margin-bottom: 0.5rem;">{desc}</div>' if desc else ''}
        """
        
        if step_res.error_message:
            html += f'<div class="error-box">{step_res.error_message}</div>'
            
        if step_res.screenshot_path:
            img_url = "/" + step_res.screenshot_path.replace("\\", "/")
            if "data/" in img_url: img_url = "/data/" + img_url.split("data/")[-1]
            html += f"""
                <div>
                    <a href="{img_url}" target="_blank">
                        <img src="{img_url}" class="screenshot-thumb" alt="Screenshot for step {step_number}">
                    </a>
                </div>
            """
            
        html += """
                    </div>
                </div>
        """
        
    html += """
            </div>
        </div>
    </body>
    </html>
    """
    return html

@router.get("/{run_id}/word")
def get_execution_word(run_id: int, db: Session = Depends(get_db)):
    try:
        import docx
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.style import WD_STYLE_TYPE
    except ImportError as e:
        raise HTTPException(status_code=500, detail=f"python-docx is not installed: {str(e)}")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Unexpected error importing python-docx: {str(e)}")
        
    run = db.query(models.ExecutionRun).filter(models.ExecutionRun.id == run_id).first()
    if not run:
        raise HTTPException(status_code=404, detail="Run not found")
        
    script = db.query(models.TestScript).filter(models.TestScript.id == run.script_id).first()
    
    document = Document()
    
    # Title
    title = document.add_heading(f'Test Execution Report: {script.name if script else "Unknown"}', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Meta Info
    document.add_paragraph(f'Run ID: {run.id}')
    document.add_paragraph(f'Status: {run.status}')
    document.add_paragraph(f'Started: {run.started_at.strftime("%Y-%m-%d %H:%M:%S") if run.started_at else "N/A"}')
    if run.completed_at and run.started_at:
        duration = f"{(run.completed_at - run.started_at).total_seconds():.1f}s"
        document.add_paragraph(f'Duration: {duration}')
        
    document.add_page_break()
    
    # Steps
    document.add_heading('Execution Steps', level=1)
    
    def translate_to_english(step, step_result):
        if not step:
            return "Raw code execution step"
            
        action = step.action.lower() if step.action else ""
        target = step.target_name
        
        if not target and step.locator_value:
            import re
            match = re.search(r'name="([^"]+)"', step.locator_value)
            if match:
                target = match.group(1)
            elif "text=" in step.locator_value:
                match = re.search(r'text="([^"]+)"', step.locator_value)
                if match:
                    target = match.group(1)
                    
        if not target:
            target = "the element"
            
        value = step.current_value
        if step.is_sensitive:
            value = "********"
            
        if action in ["fill", "input"]:
            return f"Enter data in '{target}'"
        elif action == "click":
            return f"Click on '{target}'"
        elif action == "navigate":
            return f"Navigate to application URL"
        elif action == "waitfor":
            return f"Wait for application to load"
        else:
            return f"{action.capitalize()} {target}"
    
    sorted_results = sorted(run.step_results, key=lambda s: s.step.step_number if s.step else 0)
    
    current_group = None
    group_step_counter = 1
    group_header_idx = 1
    
    for i, step_result in enumerate(sorted_results):
        step = step_result.step
        if not step: continue
        
        status_text = "Passed" if step_result.status == "Passed" else ("Failed" if step_result.status == "Failed" else "Skipped")
        desc = translate_to_english(step, step_result)
        
        group_name = step.component_group if step.component_group and step.component_group.strip() else None
        
        # Check if we are starting a new group or a standalone step
        if group_name != current_group or not group_name:
            current_group = group_name
            group_step_counter = 1
            
            # Print the header
            h = document.add_heading(level=2)
            header_text = f"Step-{group_header_idx}  {group_name}" if group_name else f"Step-{group_header_idx}  {desc}"
            r = h.add_run(header_text)
            r.font.color.rgb = docx.shared.RGBColor(21, 99, 235) # Corporate Blue
            group_header_idx += 1
            
        # Print the natural language instruction for this specific step
        p = document.add_paragraph()
        prefix = f"  {group_step_counter}. " if group_name else "1. "
        runner = p.add_run(f"{prefix}{desc}")
        runner.bold = True
        
        status_run = p.add_run(f" [{status_text}]")
        status_run.italic = True
        status_run.font.color.rgb = docx.shared.RGBColor(34, 197, 94) if status_text == "Passed" else docx.shared.RGBColor(239, 68, 68)
        
        if step_result.error_message:
            err_p = document.add_paragraph(f"  Error: {step_result.error_message}", style='Intense Quote')
            err_p.runs[0].font.color.rgb = docx.shared.RGBColor(239, 68, 68)
            
        group_step_counter += 1
        
        # Determine if this is the last step in the current group
        next_step = sorted_results[i+1].step if i + 1 < len(sorted_results) else None
        next_group = next_step.component_group if next_step and next_step.component_group and next_step.component_group.strip() else None
        
        is_last_in_group = not group_name or next_group != group_name or status_text == "Failed"
        
        # Only print the screenshot if it's the last step in the group (or standalone)
        if is_last_in_group and step_result.screenshot_path:
            import os
            img_path = step_result.screenshot_path
            if os.path.exists(img_path):
                try:
                    document.add_picture(img_path, width=Inches(6.5))
                    document.add_paragraph("")
                except Exception as e:
                    document.add_paragraph(f"(Screenshot unavailable: {str(e)})")
            
            document.add_page_break()
        
    import tempfile
    from fastapi.responses import FileResponse
    import os
    
    tmp_path = os.path.join(tempfile.gettempdir(), f"Run_{run.id}_Report.docx")
    document.save(tmp_path)
    
    return FileResponse(
        tmp_path, 
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        filename=f"Run_{run.id}_Documentation.docx"
    )
